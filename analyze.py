"""Analyze a public insurance-cost file with pandas, SQL, and Excel.

Each row is one policyholder: age, sex, bmi, children, smoker, region, charges.
"""

import math
import sqlite3
import urllib.request
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data" / "insurance.csv"
DB = ROOT / "data" / "policies.sqlite"
OUTPUT = ROOT / "output"
SOURCE = (
    "https://raw.githubusercontent.com/stedy/"
    "Machine-Learning-with-R-datasets/master/insurance.csv"
)

SUMMARY_SQL = """
SELECT smoker,
       COUNT(*) AS policies,
       ROUND(AVG(charges), 2) AS avg_charges
FROM policies
GROUP BY smoker
ORDER BY avg_charges DESC
"""

REGION_SQL = """
SELECT region,
       smoker,
       COUNT(*) AS policies,
       ROUND(AVG(charges), 2) AS avg_charges
FROM policies
GROUP BY region, smoker
ORDER BY avg_charges DESC
"""

AGE_SQL = """
SELECT CASE
           WHEN age <= 29 THEN '18-29'
           WHEN age <= 44 THEN '30-44'
           ELSE '45-64'
       END AS age_band,
       smoker,
       COUNT(*) AS policies,
       ROUND(AVG(charges), 2) AS avg_charges
FROM policies
GROUP BY age_band, smoker
ORDER BY age_band, avg_charges DESC
"""


def load_policies() -> pd.DataFrame:
    DATA.parent.mkdir(parents=True, exist_ok=True)
    if not DATA.exists():
        urllib.request.urlretrieve(SOURCE, DATA)
    frame = pd.read_csv(DATA)
    expected = {"age", "sex", "bmi", "children", "smoker", "region", "charges"}
    missing = expected - set(frame.columns)
    if missing:
        raise SystemExit(f"Unexpected columns, missing {sorted(missing)}")
    frame["smoker"] = frame["smoker"].str.lower().str.strip()
    frame["region"] = frame["region"].str.lower().str.strip()
    frame["sex"] = frame["sex"].str.lower().str.strip()
    return frame


def write_database(frame: pd.DataFrame) -> None:
    if DB.exists():
        DB.unlink()
    with sqlite3.connect(DB) as connection:
        frame.to_sql("policies", connection, index=False)


def query(sql: str) -> pd.DataFrame:
    with sqlite3.connect(DB) as connection:
        return pd.read_sql_query(sql, connection)


def write_outputs(frame: pd.DataFrame) -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    frame.to_csv(OUTPUT / "policies.csv", index=False)
    write_excel(frame, OUTPUT / "policies.xlsx")


def write_excel(frame: pd.DataFrame, path: Path) -> None:
    workbook = Workbook()
    policies = workbook.active
    policies.title = "policies"
    for row in dataframe_to_rows(frame, index=False, header=True):
        policies.append(row)
    table = Table(displayName="Policies", ref=f"A1:G{len(frame) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    policies.add_table(table)
    for column in policies.columns:
        policies.column_dimensions[column[0].column_letter].width = 14

    summary = workbook.create_sheet("summary", 0)
    summary["A1"] = "Insurance cost summary"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A3"] = "smoker"
    summary["B3"] = "policies"
    summary["C3"] = "average charges"
    for cell in ("A3", "B3", "C3"):
        summary[cell].font = Font(bold=True)
    summary["A4"] = "yes"
    summary["B4"] = '=COUNTIF(Policies[smoker],"yes")'
    summary["C4"] = '=AVERAGEIF(Policies[smoker],"yes",Policies[charges])'
    summary["A5"] = "no"
    summary["B5"] = '=COUNTIF(Policies[smoker],"no")'
    summary["C5"] = '=AVERAGEIF(Policies[smoker],"no",Policies[charges])'
    summary["A7"] = "smoker multiple"
    summary["C7"] = "=C4/C5"
    summary["C4"].number_format = '"$"#,##0.00'
    summary["C5"].number_format = '"$"#,##0.00'
    summary["C7"].number_format = "0.00"
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 14
    summary.column_dimensions["C"].width = 20

    chart = BarChart()
    chart.type = "col"
    chart.title = "Average charges by smoking status"
    chart.y_axis.title = "Average charges"
    data = Reference(summary, min_col=3, min_row=3, max_row=5)
    categories = Reference(summary, min_col=1, min_row=4, max_row=5)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.shape = 4
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = False
    chart.legend = None
    summary.add_chart(chart, "A9")

    bands = workbook.create_sheet("age bands", 1)
    bands["A1"] = "Average charges by age band"
    bands["A1"].font = Font(bold=True, size=14)
    bands["A3"] = "age band"
    bands["B3"] = "smoker"
    bands["C3"] = "policies"
    bands["D3"] = "average charges"
    for cell in ("A3", "B3", "C3", "D3"):
        bands[cell].font = Font(bold=True)
    band_rows = (("18-29", 18, 29), ("30-44", 30, 44), ("45-64", 45, 64))
    row_number = 4
    for label, low, high in band_rows:
        for smoker in ("yes", "no"):
            bands[f"A{row_number}"] = label
            bands[f"B{row_number}"] = smoker
            bands[f"C{row_number}"] = (
                f'=COUNTIFS(Policies[smoker],"{smoker}",'
                f'Policies[age],">={low}",Policies[age],"<={high}")'
            )
            bands[f"D{row_number}"] = (
                f'=AVERAGEIFS(Policies[charges],Policies[smoker],"{smoker}",'
                f'Policies[age],">={low}",Policies[age],"<={high}")'
            )
            bands[f"D{row_number}"].number_format = '"$"#,##0.00'
            row_number += 1
    bands.column_dimensions["A"].width = 22
    bands.column_dimensions["B"].width = 12
    bands.column_dimensions["C"].width = 14
    bands.column_dimensions["D"].width = 20
    workbook.save(path)


def welch_ttest(left: list[float], right: list[float]) -> tuple[float, float]:
    """Two-sample t test without assuming equal variance."""
    n1, n2 = len(left), len(right)
    mean1 = sum(left) / n1
    mean2 = sum(right) / n2
    var1 = sum((value - mean1) ** 2 for value in left) / (n1 - 1)
    var2 = sum((value - mean2) ** 2 for value in right) / (n2 - 1)
    standard_error = math.sqrt(var1 / n1 + var2 / n2)
    t_stat = (mean1 - mean2) / standard_error
    numerator = (var1 / n1 + var2 / n2) ** 2
    denominator = (var1 / n1) ** 2 / (n1 - 1) + (var2 / n2) ** 2 / (n2 - 1)
    return t_stat, numerator / denominator


def adjusted_smoking_gap(frame: pd.DataFrame) -> tuple[float, float, float]:
    """OLS of charges on smoking, age, BMI, children, sex, and region.

    Northeast is the region baseline. Returns the smoking coefficient,
    the cost of 10 additional years of age, and R-squared.
    """
    smoker = (frame["smoker"] == "yes").astype(float)
    sex_male = (frame["sex"] == "male").astype(float)
    region = pd.get_dummies(frame["region"], dtype=float).drop(columns=["northeast"])
    design = np.column_stack(
        [
            np.ones(len(frame)),
            smoker.to_numpy(),
            frame["age"].to_numpy(),
            frame["bmi"].to_numpy(),
            frame["children"].to_numpy(),
            sex_male.to_numpy(),
            region.to_numpy(),
        ]
    )
    charges = frame["charges"].to_numpy()
    coefficients, *_ = np.linalg.lstsq(design, charges, rcond=None)
    fitted = design @ coefficients
    residual = float(((charges - fitted) ** 2).sum())
    total = float(((charges - charges.mean()) ** 2).sum())
    return float(coefficients[1]), float(coefficients[2] * 10), 1 - residual / total


def main() -> None:
    frame = load_policies()
    write_database(frame)
    write_outputs(frame)
    by_smoker = query(SUMMARY_SQL)
    by_region = query(REGION_SQL)
    by_age = query(AGE_SQL)

    yes = by_smoker.loc[by_smoker["smoker"] == "yes", "avg_charges"].iloc[0]
    no = by_smoker.loc[by_smoker["smoker"] == "no", "avg_charges"].iloc[0]
    smoker_charges = frame.loc[frame["smoker"] == "yes", "charges"].tolist()
    other_charges = frame.loc[frame["smoker"] == "no", "charges"].tolist()
    t_stat, degrees = welch_ttest(smoker_charges, other_charges)
    nonsmoker_region = by_region.loc[by_region["smoker"] == "no", "avg_charges"]
    smoker_gap, age_10_years, r_squared = adjusted_smoking_gap(frame)
    smoker_median = float(frame.loc[frame["smoker"] == "yes", "charges"].median())
    other_median = float(frame.loc[frame["smoker"] == "no", "charges"].median())

    print(f"policies {len(frame)}")
    print(by_smoker.to_string(index=False))
    print(f"smoker_to_nonsmoker_ratio {yes / no:.2f}")
    print(f"welch_t {t_stat:.2f}")
    print(f"welch_df {degrees:.1f}")
    print(
        "nonsmoker_region_span "
        f"{nonsmoker_region.max() - nonsmoker_region.min():.0f}"
    )
    print(by_region.to_string(index=False))
    print(f"smoker_median {smoker_median:.2f}")
    print(f"nonsmoker_median {other_median:.2f}")
    print(f"adjusted_smoker_dollars {smoker_gap:.2f}")
    print(f"age_10_year_dollars {age_10_years:.2f}")
    print(f"model_r2 {r_squared:.3f}")
    print(by_age.to_string(index=False))
    print(f"wrote {OUTPUT / 'policies.csv'}")
    print(f"wrote {OUTPUT / 'policies.xlsx'}")


if __name__ == "__main__":
    main()
